"""Import devices from unified Excel collection template."""

from datetime import date, datetime
from io import BytesIO
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session, joinedload

from app.models import Brand, Device, DeviceModel, DeviceStatus, DeviceType, Directorate, Province, User
from app.services.labels import DEVICE_STATUS_AR, DEVICE_TYPE_AR

TEMPLATE_SHEET = "بيانات الأجهزة"

TYPE_AR_TO_ENUM = {v.replace("جهاز ", ""): k for k, v in DEVICE_TYPE_AR.items()}
TYPE_AR_TO_ENUM.update({
    "مكتبي": DeviceType.DESKTOP,
    "محمول": DeviceType.MOBILE,
    "عجلة": DeviceType.WHEEL,
    "desktop": DeviceType.DESKTOP,
    "mobile": DeviceType.MOBILE,
    "wheel": DeviceType.WHEEL,
})

STATUS_AR_TO_ENUM = {v: k for k, v in DEVICE_STATUS_AR.items()}
STATUS_AR_TO_ENUM.update({
    "working": DeviceStatus.WORKING,
    "consumed": DeviceStatus.CONSUMED,
    "يصلح للعمل": DeviceStatus.WORKING,
    "مستهلك": DeviceStatus.CONSUMED,
})

HEADER_MARKERS = ("الرقم التسلسلي", "serial")


def _normalize(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_date(value) -> Optional[date]:
    if not value:
        return None
    if hasattr(value, "date"):
        return value.date() if hasattr(value, "hour") else value
    text = _normalize(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _find_header_row(ws) -> Optional[int]:
    for row in range(1, 15):
        values = [_normalize(ws.cell(row=row, column=c).value) for c in range(1, 6)]
        joined = " ".join(values)
        if any(m in joined for m in HEADER_MARKERS):
            return row
    return None


def _build_lookups(db: Session):
    provinces = {p.name_ar: p for p in db.query(Province).all()}
    directorates = {d.name_ar: d for d in db.query(Directorate).options(joinedload(Directorate.province)).all()}
    models = (
        db.query(DeviceModel)
        .options(joinedload(DeviceModel.brand))
        .filter(DeviceModel.is_active.is_(True))
        .all()
    )
    model_lookup = {}
    for m in models:
        if m.brand:
            key = (m.brand.name_ar.strip(), m.name.strip())
            model_lookup[key] = m
            model_lookup[(m.brand.name.strip(), m.name.strip())] = m
    return provinces, directorates, model_lookup


def import_devices_from_excel(db: Session, file_content: bytes, user: User) -> dict:
    wb = load_workbook(BytesIO(file_content), data_only=True)
    if TEMPLATE_SHEET not in wb.sheetnames:
        raise ValueError(f"الورقة المطلوبة «{TEMPLATE_SHEET}» غير موجودة في الملف")

    ws = wb[TEMPLATE_SHEET]
    header_row = _find_header_row(ws)
    if not header_row:
        raise ValueError("تعذّر العثور على صف العناوين في الملف")

    headers = {_normalize(ws.cell(row=header_row, column=c).value): c for c in range(1, 25)}

    def col(name_part: str) -> Optional[int]:
        for h, idx in headers.items():
            if name_part in h:
                return idx
        return None

    col_serial = col("التسلسلي") or col("serial")
    col_mfg = col("المصنعي") or col("manufacturer")
    col_asset = col("الأصل")
    col_type = col("نوع الجهاز") or col("نوع")
    col_brand = col("الشركة")
    col_model = col("الموديل")
    col_province = col("المحافظة")
    col_directorate = col("المديرية")
    col_status = col("حالة")
    col_workplace = col("مكان العمل")
    col_location = col("الموقع")
    col_department = col("القسم")
    col_assigned = col("المسؤول")
    col_condition = col("وصف الحالة")
    col_purchase = col("تاريخ الشراء")
    col_received = col("تاريخ الاستلام")
    col_warranty = col("انتهاء الضمان")
    col_notes = col("ملاحظات")

    if not col_serial:
        raise ValueError("عمود الرقم التسلسلي غير موجود")

    provinces, directorates, model_lookup = _build_lookups(db)

    created = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    for row in range(header_row + 1, ws.max_row + 1):
        serial = _normalize(ws.cell(row=row, column=col_serial).value)
        if not serial:
            continue
        if "مثال" in serial or serial.lower() == "serial":
            skipped += 1
            continue

        notes_val = _normalize(ws.cell(row=row, column=col_notes).value) if col_notes else ""
        if "مثال" in notes_val and "احذف" in notes_val:
            skipped += 1
            continue

        try:
            brand_name = _normalize(ws.cell(row=row, column=col_brand).value) if col_brand else ""
            model_name = _normalize(ws.cell(row=row, column=col_model).value) if col_model else ""
            province_name = _normalize(ws.cell(row=row, column=col_province).value) if col_province else ""
            directorate_name = _normalize(ws.cell(row=row, column=col_directorate).value) if col_directorate else ""
            type_raw = _normalize(ws.cell(row=row, column=col_type).value) if col_type else "محمول"
            status_raw = _normalize(ws.cell(row=row, column=col_status).value) if col_status else "يصلح للعمل"

            device_type = TYPE_AR_TO_ENUM.get(type_raw)
            if not device_type:
                raise ValueError(f"نوع جهاز غير معروف: {type_raw}")

            status = STATUS_AR_TO_ENUM.get(status_raw)
            if not status:
                for key, val in STATUS_AR_TO_ENUM.items():
                    if key in status_raw or status_raw in key:
                        status = val
                        break
            if not status:
                raise ValueError(f"حالة غير معروفة: {status_raw}")

            province = provinces.get(province_name)
            if not province:
                raise ValueError(f"محافظة غير معروفة: {province_name}")

            directorate = directorates.get(directorate_name)
            if not directorate:
                raise ValueError(f"مديرية غير معروفة: {directorate_name}")

            model = model_lookup.get((brand_name, model_name))
            if not model:
                raise ValueError(f"موديل غير مسجّل: {brand_name} — {model_name}")

            mfg_serial = _normalize(ws.cell(row=row, column=col_mfg).value) if col_mfg else None
            workplace = _normalize(ws.cell(row=row, column=col_workplace).value) if col_workplace else None

            if not mfg_serial:
                raise ValueError("الرقم المصنعي مطلوب")
            if not workplace:
                raise ValueError("مكان العمل مطلوب")

            payload = {
                "manufacturer_serial": mfg_serial,
                "asset_number": _normalize(ws.cell(row=row, column=col_asset).value) or None if col_asset else None,
                "model_id": model.id,
                "province_id": province.id,
                "directorate_id": directorate.id,
                "status": status,
                "device_type": device_type,
                "workplace": workplace,
                "location": _normalize(ws.cell(row=row, column=col_location).value) or None if col_location else None,
                "department": _normalize(ws.cell(row=row, column=col_department).value) or None if col_department else None,
                "assigned_to": _normalize(ws.cell(row=row, column=col_assigned).value) or None if col_assigned else None,
                "condition_notes": _normalize(ws.cell(row=row, column=col_condition).value) or None if col_condition else None,
                "purchase_date": _parse_date(ws.cell(row=row, column=col_purchase).value) if col_purchase else None,
                "received_date": _parse_date(ws.cell(row=row, column=col_received).value) if col_received else None,
                "warranty_expiry": _parse_date(ws.cell(row=row, column=col_warranty).value) if col_warranty else None,
                "notes": notes_val or None,
            }

            existing = db.query(Device).filter(Device.serial_number == serial).first()
            if existing:
                for key, val in payload.items():
                    setattr(existing, key, val)
                updated += 1
            else:
                db.add(Device(serial_number=serial, created_by_id=user.id, **payload))
                created += 1

        except Exception as exc:
            errors.append(f"سطر {row} ({serial}): {exc}")

    if created or updated:
        db.commit()

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "total_processed": created + updated,
    }
