from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.models import BookDeviceItem, Device, DeviceModel, OfficialBook, Province
from app.services.labels import BOOK_TYPE_AR, DEVICE_STATUS_AR, DEVICE_TYPE_AR


def _style_header(ws, row: int, col_count: int):
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def export_devices_excel(
    db: Session,
    province_id: Optional[int] = None,
    status: Optional[str] = None,
    device_type: Optional[str] = None,
) -> BytesIO:
    query = (
        db.query(Device)
        .options(joinedload(Device.model).joinedload(DeviceModel.brand))
        .options(joinedload(Device.province))
    )
    if province_id:
        query = query.filter(Device.province_id == province_id)
    if status:
        query = query.filter(Device.status == status)
    if device_type:
        query = query.filter(Device.device_type == device_type)

    devices: List[Device] = query.order_by(Device.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "الأجهزة"
    ws.sheet_view.rightToLeft = True

    headers = [
        "م",
        "الرقم التسلسلي",
        "رقم الأصل",
        "نوع الجهاز",
        "الشركة",
        "الموديل",
        "المحافظة",
        "الحالة",
        "الموقع",
        "تاريخ الشراء",
        "ملاحظات",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for idx, d in enumerate(devices, 1):
        ws.append([
            idx,
            d.serial_number,
            d.asset_number or "",
            DEVICE_TYPE_AR.get(d.device_type, str(d.device_type)),
            d.model.brand.name_ar if d.model and d.model.brand else "",
            d.model.name if d.model else "",
            d.province.name_ar if d.province else "",
            DEVICE_STATUS_AR.get(d.status, str(d.status)),
            d.location or "",
            str(d.purchase_date) if d.purchase_date else "",
            d.notes or "",
        ])

    _auto_width(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_books_excel(
    db: Session,
    province_id: Optional[int] = None,
    book_type: Optional[str] = None,
) -> BytesIO:
    query = (
        db.query(OfficialBook)
        .options(joinedload(OfficialBook.province))
        .options(
            joinedload(OfficialBook.device_items)
            .joinedload(BookDeviceItem.device)
            .joinedload(Device.model)
            .joinedload(DeviceModel.brand)
        )
    )
    if province_id:
        query = query.filter(OfficialBook.province_id == province_id)
    if book_type:
        query = query.filter(OfficialBook.book_type == book_type)

    books: List[OfficialBook] = query.order_by(OfficialBook.book_date.desc()).all()

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "الكتب الرسمية"
    ws_summary.sheet_view.rightToLeft = True

    summary_headers = [
        "م",
        "رقم الكتاب",
        "النوع",
        "التاريخ",
        "المحافظة",
        "الموضوع",
        "من",
        "إلى",
        "عدد الأجهزة",
        "ملاحظات",
    ]
    ws_summary.append(summary_headers)
    _style_header(ws_summary, 1, len(summary_headers))

    ws_details = wb.create_sheet("تفاصيل الأجهزة")
    ws_details.sheet_view.rightToLeft = True
    detail_headers = [
        "رقم الكتاب",
        "الرقم التسلسلي",
        "نوع الجهاز",
        "الشركة",
        "الموديل",
        "الكمية",
        "ملاحظات",
    ]
    ws_details.append(detail_headers)
    _style_header(ws_details, 1, len(detail_headers))

    for idx, book in enumerate(books, 1):
        ws_summary.append([
            idx,
            book.book_number,
            BOOK_TYPE_AR.get(book.book_type, str(book.book_type)),
            str(book.book_date),
            book.province.name_ar if book.province else "",
            book.subject,
            book.from_entity,
            book.to_entity,
            len(book.device_items),
            book.notes or "",
        ])
        for item in book.device_items:
            device = item.device
            ws_details.append([
                book.book_number,
                device.serial_number if device else "",
                DEVICE_TYPE_AR.get(device.device_type, "") if device else "",
                device.model.brand.name_ar if device and device.model and device.model.brand else "",
                device.model.name if device and device.model else "",
                item.quantity,
                item.notes or "",
            ])

    _auto_width(ws_summary)
    _auto_width(ws_details)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
