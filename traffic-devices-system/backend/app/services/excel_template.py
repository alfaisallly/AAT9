"""Generate unified Excel template for province/HQ device data collection."""

from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session, joinedload

from app.models import Brand, DeviceModel, DeviceStatus, DeviceType, Directorate, Province
from app.services.labels import DEVICE_STATUS_AR, DEVICE_TYPE_AR

TEMPLATE_SHEET = "بيانات الأجهزة"
HEADER_ROW = 6
FIRST_DATA_ROW = 9
MAX_DATA_ROWS = 502

COLUMNS = [
    ("م", 6, False),
    ("الرقم التسلسلي الداخلي *", 22, True),
    ("الرقم المصنعي (تسلسل التصنيع) *", 26, True),
    ("رقم الأصل", 16, False),
    ("نوع الجهاز *", 14, True),
    ("الشركة المصنعة *", 18, True),
    ("الموديل *", 20, True),
    ("المحافظة *", 16, True),
    ("المديرية *", 28, True),
    ("حالة الجهاز *", 22, True),
    ("مكان العمل *", 22, True),
    ("الموقع التفصيلي", 22, False),
    ("القسم / الوحدة", 18, False),
    ("المسؤول / المستخدم", 18, False),
    ("وصف الحالة / العطل", 24, False),
    ("تاريخ الشراء", 14, False),
    ("تاريخ الاستلام", 14, False),
    ("انتهاء الضمان", 14, False),
    ("ملاحظات", 24, False),
]

EXAMPLE_ROWS = [
    [
        1,
        "HT-BGD-0001",
        "MSN-HY-78451236",
        "AM-14001",
        "محمول",
        "هايتيرا",
        "HP780",
        "بغداد",
        "مديرية مرور بغداد - الكرخ",
        "يصلح للعمل",
        "شعبة اتصالات الكرخ",
        "مقر مديرية المرور - الكرخ",
        "قسم الاتصالات",
        "ملازم أحمد",
        "",
        "2022-03-15",
        "2022-04-01",
        "2025-04-01",
        "مثال — احذف هذا السطر قبل الإرسال",
    ],
    [
        2,
        "SP-BSR-0002",
        "MSN-SEP-99887766",
        "",
        "عجلة",
        "سيبورا",
        "SRG3900",
        "البصرة",
        "مديرية مرور البصرة",
        "مستهلك - لا يصلح للعمل",
        "مركز مرور الزبير",
        "مديرية مرور البصرة",
        "وحدة الاتصالات",
        "رئيس مركز",
        "عطل لوحة المفاتيح",
        "2019-06-10",
        "2019-07-01",
        "",
        "مثال — احذف هذا السطر قبل الإرسال",
    ],
]

HEADER_FILL = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
REQUIRED_FILL = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
EXAMPLE_FILL = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
META_FILL = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _style_header_row(ws, row: int, col_count: int):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        if COLUMNS[col - 1][2]:
            cell.fill = PatternFill(start_color="154360", end_color="154360", fill_type="solid")


def _add_list_validation(ws, col_letter: str, formula: str, start_row: int = FIRST_DATA_ROW):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "اختر قيمة من القائمة المعتمدة"
    dv.errorTitle = "قيمة غير صالحة"
    dv.prompt = "اختر من القائمة"
    dv.promptTitle = "قائمة معتمدة"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{start_row}:{col_letter}{MAX_DATA_ROWS}")


def _write_reference_sheet(ws, title: str, headers: list[str], rows: list[list]):
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col_idx, width in enumerate([max(len(h), 12) for h in headers], 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 4, 40)


def generate_devices_import_template(db: Session, province_hint: Optional[str] = None) -> BytesIO:
    provinces = db.query(Province).order_by(Province.name_ar).all()
    directorates = (
        db.query(Directorate)
        .options(joinedload(Directorate.province))
        .filter(Directorate.is_active.is_(True))
        .order_by(Directorate.name_ar)
        .all()
    )
    brands = db.query(Brand).filter(Brand.is_active.is_(True)).order_by(Brand.name_ar).all()
    models = (
        db.query(DeviceModel)
        .options(joinedload(DeviceModel.brand))
        .filter(DeviceModel.is_active.is_(True))
        .order_by(DeviceModel.name)
        .all()
    )

    wb = Workbook()

    # ── Instructions ──
    ws_info = wb.active
    ws_info.title = "التعليمات"
    ws_info.sheet_view.rightToLeft = True
    ws_info.column_dimensions["A"].width = 100
    instructions = [
        "نموذج جمع بيانات الأجهزة اللاسلكية — مديرية المرور العامة",
        "",
        "■ الغرض: تعبئة بيانات الأجهزة في المحافظات والمقر المركزي لرفعها إلى النظام.",
        "",
        "■ خطوات التعبئة:",
        "  1. املأ بيانات الجهة في أعلى ورقة «بيانات الأجهزة» (المحافظة، المديرية، المسؤول...).",
        "  2. احذف سطور الأمثلة الصفراء قبل إدخال بياناتكم.",
        "  3. أدخل كل جهاز في سطر — الحقول التي بها * إلزامية.",
        "  4. استخدم القوائم المنسدلة للمحافظة، المديرية، النوع، الشركة، والحالة.",
        "  5. التواريخ بصيغة: YYYY-MM-DD  (مثال: 2024-05-15).",
        "",
        "■ الحقول الإلزامية:",
        "  • الرقم التسلسلي الداخلي: رقم فريد لكل جهاز داخل مديريتكم.",
        "  • الرقم المصنعي: تسلسل التصنيع من الشركة المصنعة (يُستخدم في البحث).",
        "  • نوع الجهاز: مكتبي / محمول / عجلة.",
        "  • الشركة والموديل: من القائمة أو ورقة «الشركات والموديلات».",
        "  • المحافظة والمديرية: من القوائم المرجعية.",
        "  • حالة الجهاز: يصلح للعمل / مستهلك - لا يصلح للعمل.",
        "  • مكان العمل: موقع العمل الفعلي للجهاز.",
        "",
        "■ أوراق مرجعية:",
        "  • المحافظات — قائمة المحافظات المعتمدة.",
        "  • المديريات — المديريات والمحافظات التابعة لها.",
        "  • الشركات والموديلات — الشركات والموديلات المسجّلة في النظام.",
        "  • القيم المعتمدة — أنواع الأجهزة والحالات.",
        "",
        "■ بعد التعبئة: أرسل الملف إلى المقر المركزي أو ارفعه عبر النظام (استيراد Excel).",
    ]
    for i, line in enumerate(instructions, 1):
        cell = ws_info.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        if i == 1:
            cell.font = Font(bold=True, size=14, color="1A5276")

    # ── Reference sheets (create before main for validation refs) ──
    ws_types = wb.create_sheet("القيم المعتمدة")
    ws_types.sheet_view.rightToLeft = True
    ws_types.append(["نوع الجهاز", "الحالة"])
    type_values = [DEVICE_TYPE_AR[t].replace("جهاز ", "") for t in DeviceType]
    status_values = [DEVICE_STATUS_AR[s] for s in DeviceStatus]
    for i in range(max(len(type_values), len(status_values))):
        ws_types.append([
            type_values[i] if i < len(type_values) else "",
            status_values[i] if i < len(status_values) else "",
        ])
    ws_types.column_dimensions["A"].width = 18
    ws_types.column_dimensions["B"].width = 28

    ws_prov = wb.create_sheet("المحافظات")
    ws_prov.sheet_view.rightToLeft = True
    ws_prov.append(["المحافظة", "الرمز"])
    for p in provinces:
        ws_prov.append([p.name_ar, p.code])
    ws_prov.column_dimensions["A"].width = 22
    ws_prov.column_dimensions["B"].width = 10

    ws_dir = wb.create_sheet("المديريات")
    ws_dir.sheet_view.rightToLeft = True
    ws_dir.append(["المديرية", "المحافظة", "الرمز"])
    for d in directorates:
        prov_name = d.province.name_ar if d.province else "—"
        ws_dir.append([d.name_ar, prov_name, d.code])
    ws_dir.column_dimensions["A"].width = 36
    ws_dir.column_dimensions["B"].width = 18
    ws_dir.column_dimensions["C"].width = 10

    ws_models = wb.create_sheet("الشركات والموديلات")
    ws_models.sheet_view.rightToLeft = True
    ws_models.append(["الشركة", "الموديل", "نوع الجهاز"])
    for m in models:
        brand_name = m.brand.name_ar if m.brand else ""
        type_label = DEVICE_TYPE_AR.get(m.device_type, str(m.device_type)).replace("جهاز ", "")
        ws_models.append([brand_name, m.name, type_label])
    ws_models.column_dimensions["A"].width = 16
    ws_models.column_dimensions["B"].width = 22
    ws_models.column_dimensions["C"].width = 14

    # ── Main data sheet ──
    ws = wb.create_sheet(TEMPLATE_SHEET, 0)
    ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:S1")
    title_cell = ws["A1"]
    title_cell.value = "نموذج جمع بيانات الأجهزة اللاسلكية — مديرية المرور العامة"
    title_cell.font = Font(bold=True, size=14, color="1A5276")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = META_FILL

    meta_labels = [
        ("A2", "المحافظة:"),
        ("D2", "المديرية:"),
        ("H2", "تاريخ التعبئة:"),
        ("K2", "اسم المسؤول:"),
        ("A3", "الجهة / القسم المرسل:"),
        ("H3", "رقم الهاتف:"),
        ("K3", "البريد الإلكتروني:"),
    ]
    for cell_ref, label in meta_labels:
        c = ws[cell_ref]
        c.value = label
        c.font = Font(bold=True, size=10, color="475569")
        c.fill = META_FILL
    ws.merge_cells("B2:C2")
    ws.merge_cells("E2:G2")
    ws.merge_cells("I2:J2")
    ws.merge_cells("L2:S2")
    ws.merge_cells("B3:G3")
    ws.merge_cells("I3:J3")
    ws.merge_cells("L3:S3")

    if province_hint:
        ws["B2"] = province_hint

    ws["A5"] = "▼ أدخل بيانات الأجهزة من السطر 9 فما فوق — احذف سطور الأمثلة الصفراء"
    ws["A5"].font = Font(italic=True, color="7D6608", size=10)
    ws.merge_cells("A5:S5")

    for col_idx, (header, width, required) in enumerate(COLUMNS, 1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    _style_header_row(ws, HEADER_ROW, len(COLUMNS))

    for row_offset, example in enumerate(EXAMPLE_ROWS):
        row_num = HEADER_ROW + 1 + row_offset
        for col_idx, value in enumerate(example, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.fill = EXAMPLE_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "right", vertical="center")

    ws.row_dimensions[HEADER_ROW].height = 32
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"

    last_prov = len(provinces) + 1
    last_dir = len(directorates) + 1
    last_brand = len(brands) + 1
    last_type = len(type_values) + 1
    last_status = len(status_values) + 1

    _add_list_validation(ws, "E", f"='القيم المعتمدة'!$A$2:$A${last_type}")
    _add_list_validation(ws, "F", f"='الشركات والموديلات'!$A$2:$A${max(last_brand, len(models) + 1)}")
    _add_list_validation(ws, "H", f"='المحافظات'!$A$2:$A${last_prov}")
    _add_list_validation(ws, "I", f"='المديريات'!$A$2:$A${last_dir}")
    _add_list_validation(ws, "J", f"='القيم المعتمدة'!$B$2:$B${last_status}")

    for row in range(FIRST_DATA_ROW, MAX_DATA_ROWS + 1):
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = THIN_BORDER
            if col > 1:
                cell.alignment = Alignment(horizontal="right", vertical="center")

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
